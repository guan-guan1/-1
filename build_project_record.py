from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out=Path(r'D:\chatgpt\codex\打卡\项目工作记录.docx')
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=doc.styles
for nm in ['Normal','Title','Heading 1','Heading 2']:
    st=styles[nm]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor(45,55,65)
styles['Title'].font.size=Pt(24); styles['Title'].font.bold=True; styles['Title'].font.color.rgb=RGBColor(31,57,78)
styles['Heading 1'].font.size=Pt(15); styles['Heading 1'].font.bold=True; styles['Heading 1'].font.color.rgb=RGBColor(31,94,112)
styles['Heading 2'].font.size=Pt(11); styles['Heading 2'].font.bold=True

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell(cell, text, bold=False, color=None, size=8.2):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
    r=p.add_run(str(text)); r.bold=bold; r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    t.autofit=False
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i],h,True,(255,255,255),8.5); shade(t.rows[0].cells[i],'1F5E70')
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            set_cell(cells[i],v,size=8.1)
            if len(t.rows)%2==0: shade(cells[i],'F1F6F7')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def para(text):
    p=doc.add_paragraph(text); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.08; return p

p=doc.add_paragraph(style='Title'); p.add_run('打卡项目工作记录')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(8)
r=p.add_run('记录日期：2026年10月27日    |    版本：1.0'); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(90,105,115)
para('本记录汇总截至记录日期可查的项目沟通、当前工作区文件及 Git 提交记录。项目目标是提供一个面向 Android 的 PWA，用于记录是否遵守晚间不点外卖的约束，并支持多个打卡计划、日历查看与跨设备同步。以下将已确认事实、用户报告和待核实事项分开标注。')

doc.add_heading('一、项目概况与目标', level=1)
para('产品形态：可安装到 Android 主屏幕的渐进式 Web 应用（PWA），当前通过 GitHub Pages 提供网页。核心交互是打开后看到当前打卡计划、计划切换栏和小型日历；用户可在规定时间段内标记成功或失败，并查看每天备注。')
para('项目边界：目前记录由用户主动完成；不读取支付记录，也不自动判断是否在淘宝闪购或美团消费。后续长期记录保存和跨设备同步目标由云端服务承担。')
para('项目启动的准确日期：信息缺失。Git 仓库可确认的最早提交为 2026年10月24日（Initial commit）。')

doc.add_heading('二、已完成工作及产出物', level=1)
table(['阶段 / 时间','工作与产出','状态与依据'],[
('产品定义（日期信息不完整）','从逐日习惯回顾调整为“正在进行的打卡计划 + 小型日历”；支持多计划、时间段、成功/失败记录和备注。','方向由用户在沟通中确认；具体需求细节见 README。'),
('基础应用（2026-10-24）','建立应用仓库和 PWA 页面，包含计划展示、打卡操作和日历等基础能力。','Git：c2c60fa Initial commit；后续功能以 README 与代码为依据。'),
('打卡规则迭代','增加失败按钮“嘴馋了，没忍住”；不同计划设置各自时间段；跨午夜时段归属前一天；已记录过去日期只读、无记录日期可补录；编辑/删除计划、备注与失败改成功确认。','需求由用户逐项提出并确认，当前版本功能描述见 README；不是逐项独立验收报告。'),
('PWA 与 GitHub Pages','配置 PWA 安装资源、离线缓存及 Pages Actions 发布流程。','工作区包含 public、service-worker.js 与 .github/workflows/pages.yml；当前工作流改动尚未提交。'),
('Supabase 云同步（2026-10-24）','接入邮箱登录及云端同步；迭代登录/注册、密码恢复、跨设备删除同步与刷新循环修复。','Git：644a6d1、e5e489a、390b353、4ff70ce、801d57c。'),
('连续天数与重试提示（2026-10-27）','增加连续成功打卡天数显示及云同步重试提示。','Git：4093ec5。用户报告此前跨设备同步修复后暂未再观察到可见故障；属用户测试反馈。'),
('迁移到腾讯云 CloudBase（当前工作区）','新增 CloudBase 客户端同步、项目配置、数据规则文件与 Vite 构建配置；为国内访问及减少服务器运维负担进行迁移。','代码在工作区但未提交；本地 Vite build 曾成功。GitHub Actions 发布及实际云端联调尚未确认。')
], [1.15,3.15,2.65])

doc.add_heading('三、关键决策与变更记录', level=1)
table(['事项','已确认决定','依据 / 影响'],[
('产品定位','记录遵守晚间不点外卖的约束，而非培养习惯或提供逐日习惯回顾。','用户明确修正产品方向。'),
('记录方式','用户主动点击成功或失败；不读取支付流水、不自动判定外卖消费。','MVP 范围与隐私边界。'),
('跨日与历史记录','跨午夜时段归到前一天；过去日期已有记录不可改，无记录可补录。','用户明确确认。'),
('云端方向','先采用 Supabase，后因境内连接不稳定及运维偏好转向腾讯云 CloudBase 上海免费环境。','CloudBase 环境 ID：my-daka-d7g92edvi8d023606；区域 ap-shanghai，由用户提供。'),
('网站托管与访问域名','GitHub Pages 地址继续作为应用访问入口；CloudBase 静态托管默认域名与 GitHub Pages 来源地址用途不同。','当前 Pages 地址为 https://guan-guan1.github.io/-1/；访问安全来源配置仍待核实。'),
('旧数据迁移','代码尝试从旧 Supabase 导入，但需浏览器仍有有效 Supabase 会话且网络可达。','迁移成功与否尚未确认，不能视为已完成。')
], [1.2,3.0,2.75])

doc.add_heading('四、当前进展与未完成事项', level=1)
para('已知工作区状态：CloudBase 迁移相关文件及 Pages 工作流、README、应用入口和 service worker 均有未提交修改。未提交代码不等于已部署到线上。当前 Pages 线上版本和最近一次 Actions 发布结果，本记录未进行实时核验。')
para('已知构建情况：工作区曾成功执行本地 Vite production build；构建输出提示 JavaScript 包体超过 500 KB。构建成功不等于 CloudBase 登录、同步及迁移已经端到端验证。')
table(['事项','当前状态','责任方'],[
('CloudBase 环境认证与邮件登录配置','待确认控制台是否已开启；没有可核验的控制台操作记录。','用户'),
('数据集合及权限规则','代码库提供 user_checkin_data 规则文件；需确认控制台已创建集合并应用规则。','用户'),
('来源域名配置','需确认将 guan-guan1.github.io 加入 CloudBase 允许来源；设置入口及生效状态待确认。','用户'),
('CloudBase 迁移代码提交与发布','工作区未提交；CI/Pages 发布未确认。','助手执行；需沿用当前 GitHub 授权/连接状态'),
('端到端验证','邮箱注册/登录、同步、删除、多设备一致性、离线行为、旧 Supabase 数据迁移均需验证。','用户与助手'),
('旧数据保全','迁移前后数据数量与完整性未确认。','用户与助手')
],[1.9,3.65,1.4])

doc.add_heading('五、问题、风险及待确认事项', level=1)
table(['类别','问题或风险','状态'],[
('访问配置','CloudBase 的来源域名、安全规则和邮箱认证配置若缺失，可能造成登录/读写失败。','待确认控制台设置。'),
('部署状态','CloudBase 改动未提交，不能据此认定 GitHub Pages 已使用新后端。','由 Git 状态确认。'),
('数据迁移','旧 Supabase 导入依赖有效会话与网络；迁移遗漏或重复的结果未检查。','待测试并备份。'),
('兼容性','本地构建成功；生产发布及 Android 手机上的实际登录、同步与安装后运行未验证。','待验证。'),
('包体积','构建器提示一个 JS 包超过 500 KB，可能影响首次加载。','优化建议，非当前阻塞项。'),
('项目管理材料','未发现独立会议纪要或正式任务清单文件；项目启动日期也无法从当前证据准确确定。','信息缺失。')
],[1.1,4.4,1.45])

doc.add_heading('六、下一步计划、负责人及预计完成时间', level=1)
para('下表时间为工作量估计，不是已承诺的日历日期；实际时间取决于控制台配置、部署状态和网络情况。')
table(['顺序 / 待办','负责人','预计耗时','完成标准'],[
('1. 确认 CloudBase 邮箱登录方式及邮件发送配置','用户','约 5–10 分钟','控制台已启用对应认证方式，能发出验证邮件。'),
('2. 创建 user_checkin_data 集合并应用用户隔离读写规则','用户','约 5–10 分钟','用户只能读取和修改自己的数据。'),
('3. 配置 GitHub Pages 来源域名 guan-guan1.github.io','用户','约 5 分钟，另计配置生效时间','CloudBase 来源设置接受 Pages 请求。'),
('4. 检查 CloudBase 改动、提交并发布 Pages','助手','约 15–30 分钟','构建和 Actions 发布成功，线上代码为迁移版本。'),
('5. 多端端到端测试：注册/登录、增删计划、打卡、备注、连续天数及断网重试','用户与助手','约 20–40 分钟','电脑与 Android 端数据一致，无重复刷新或无法恢复的错误。'),
('6. 核验旧 Supabase 数据迁移并保留备份','用户与助手','视数据量及网络而定','逐项核对旧数据；迁移完成前保留原服务数据。')
],[2.8,1.0,1.1,2.0])

doc.add_heading('七、需要补充的信息', level=1)
for x in [
'项目正式启动日期；当前最早可核实的仓库提交日期为 2026年10月24日。',
'CloudBase 控制台中邮箱认证、邮件发送、集合、权限规则和来源域名的实际配置状态。',
'GitHub Actions 最新部署状态及当前线上页面是否已包含 CloudBase 迁移。',
'CloudBase 上线后，电脑端和 Android 端登录、同步与删除同步测试结果。',
'旧 Supabase 数据是否已导入 CloudBase，以及导入前后核对和备份结果。',
'如有独立会议纪要、任务清单或验收记录，请补充其文件或链接。'
]:
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(x)

# footer
for s in doc.sections:
    f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=f.add_run('打卡项目工作记录  |  2026-10-27  |  v1.0'); rr.font.name='Microsoft YaHei'; rr._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(120,130,138)

doc.save(out)
print(out)

